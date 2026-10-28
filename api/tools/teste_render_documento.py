"""Testes do modo "Formatado" do .docx (documento.html) — sem banco, sem servidor.

Rodar da pasta api/:  python tools/teste_render_documento.py

Cobre as três promessas da feature:
  1. o extrator de texto agora grava a tabela do Word no mesmo markdown da planilha
     (_linha_tabela), com o "|" das células trocado por "/";
  2. renderizar_blocos_docx devolve títulos, listas (agrupadas), runs com negrito/itálico e
     tabelas sem a repetição da mesclagem horizontal;
  3. as páginas do render são as MESMAS do texto extraído — uma por marcador [[página N]],
     na mesma ordem — porque é isso que faz o "Página 7" do Formatado ser o da IA.
"""

import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_BREAK  # noqa: E402

from funcoes.documentos_utils import (  # noqa: E402
    extrair_texto_documento,
    paginas_do_texto,
    renderizar_blocos_docx,
    texto_paginado,
)


def _documento_de_teste() -> Path:
    d = Document()
    d.add_heading("Relatório", level=0)          # estilo "Title" → h1
    d.add_heading("1. Introdução", level=1)      # Heading 1 → h1
    p = d.add_paragraph("Normal com ")
    p.add_run("negrito").bold = True
    p.add_run(" e ")
    p.add_run("itálico").italic = True
    p.add_run(" no fim.")
    d.add_paragraph("um", style="List Bullet")   # numPr herdado do estilo
    d.add_paragraph("dois", style="List Bullet")
    d.add_paragraph("primeiro", style="List Number")
    d.add_paragraph("segundo", style="List Number")
    tabela = d.add_table(rows=2, cols=3)
    tabela.alignment = WD_TABLE_ALIGNMENT.LEFT
    for j, texto in enumerate(("Nome", "CPF", "Valor|R$")):
        tabela.rows[0].cells[j].text = texto
    for j, texto in enumerate(("Fulano", "123", "10,00")):
        tabela.rows[1].cells[j].text = texto
    quebra = d.add_paragraph()
    quebra.add_run().add_break(WD_BREAK.PAGE)
    d.add_heading("Segunda página", level=1)
    d.add_paragraph("Conteúdo.")
    caminho = Path(tempfile.gettempdir()) / "analyser_teste_render.docx"
    d.save(str(caminho))
    return caminho


def main() -> int:
    caminho = _documento_de_teste()
    falhas = []

    # 1) texto extraído: tabela em markdown, célula com "|" aparada
    texto, erro = extrair_texto_documento(caminho)
    if erro:
        falhas.append(f"extração falhou: {erro}")
    else:
        if "| Nome | CPF | Valor/R$ |" not in texto:
            falhas.append("linha de cabeçalho da tabela não saiu em markdown")
        if "| --- | --- | --- |" not in texto:
            falhas.append("separador de cabeçalho (---) ausente")
        if "\t" in texto:
            falhas.append("tabela ainda vazou com tabuladores")

    # 2) blocos do render
    paginas = renderizar_blocos_docx(caminho)
    blocos = paginas[0]["blocos"]
    tipos = [b["t"] for b in blocos]
    if tipos[:2] != ["h", "h"]:
        falhas.append(f"títulos não saíram como h: {tipos[:2]}")
    formatado = next((b for b in blocos if b["t"] == "p" and "runs" in b), None)
    if not formatado:
        falhas.append("parágrafo com negrito/itálico não saiu em runs")
    else:
        band = {(bool(r["b"]), bool(r["i"])) for r in formatado["runs"]}
        if (True, False) not in band or (False, True) not in band:
            falhas.append("runs sem as marcas de negrito/itálico")
    listas = [b for b in blocos if b["t"] == "l"]
    if [len(l["itens"]) for l in listas] != [2, 2] or [l["ordem"] for l in listas] != [False, True]:
        falhas.append(f"listas não agruparam/bancaram a ordenada: {listas}")
    tabela = next((b for b in blocos if b["t"] == "t"), None)
    if not tabela or tabela["cabecalho"] != ["Nome", "CPF", "Valor|R$"]:
        falhas.append(f"tabela do render errada: {tabela}")

    # mesclagem horizontal não duplica célula
    d2 = Document()
    t2 = d2.add_table(rows=1, cols=2)
    t2.rows[0].cells[0].merge(t2.rows[0].cells[1])
    t2.rows[0].cells[0].text = "única"
    m2 = Path(tempfile.gettempdir()) / "analyser_teste_merge.docx"
    d2.save(str(m2))
    mesclada = renderizar_blocos_docx(m2)[0]["blocos"][0]
    if mesclada["linhas"] != [["única"]] and mesclada["cabecalho"] != ["única"]:
        falhas.append(f"célula mesclada duplicou: {mesclada}")

    # 3) as páginas do render são as páginas do texto
    if texto and texto_paginado(texto):
        fatiadas = paginas_do_texto(texto)
        if len(fatiadas) != len(paginas):
            falhas.append(f"render com {len(paginas)} páginas, texto com {len(fatiadas)}")

    if falhas:
        for f in falhas:
            print("FALHOU:", f)
        return 1
    print("ok: extração em markdown, blocos do render e paginação alinhada — tudo confere")
    return 0


if __name__ == "__main__":
    sys.exit(main())
