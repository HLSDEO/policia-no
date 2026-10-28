"""Relatório em Word (.docx) a partir de uma resposta do chat de IA.

Quando o investigador pede um relatório no chat — no chat do inquérito (web/ia.html) ou no
"Prompt de IA específico" (mesma tela, conversa restrita a um documento) — a resposta continua
aparecendo na conversa em markdown e, além disso, vira
um .docx pronto para baixar, com o mesmo cabeçalho/rodapé dos relatórios do sistema.

São três peças:

  - `pedido_de_relatorio(pergunta)`: reconhece o pedido no texto do investigador (verbo de
    geração perto de "relatório"/"word"/"docx"). O mesmo gatilho injeta `INSTRUCAO_RELATORIO` no
    prompt de sistema — sem isso o modelo responde em tom de conversa e o Word sai com cara de
    chat, não de documento.
  - `montar_docx(...)`: converte o markdown da resposta em .docx (títulos, listas, tabelas,
    negrito/itálico/código, citações). É o mesmo subconjunto de markdown que o renderizador do
    front entende (web/js/ia.js), que é o que os modelos de fato usam.
  - `anexar_relatorio(...)`: grava o arquivo em PASTA_INVESTIGACAO/relatorios_ia, registra em
    IA_TB05_RELATORIOS_DOCX e devolve a resposta com o link de download no fim. Nunca levanta:
    a resposta já está pronta, e falhar ao montar o Word não pode derrubar o turno.

O binário não vai para o banco (megabytes que o usuário baixa uma vez) — mesmo desenho de
IA_TB04_ARQUIVOS_PROMPT e de REL_TB04_GERACOES (funcoes/relatorios_geracao.py).
"""

import logging
import re
import unicodedata
from datetime import datetime, timedelta
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

from funcoes.db_utils import get_connection
from funcoes.investigacao_utils import PASTA_DADOS
# privadas de propósito: o cabeçalho (brasão + SERVIÇO PÚBLICO FEDERAL / MJSP - POLÍCIA
# FEDERAL), o rodapé com número de página e o sombreado de célula são os mesmos do relatório
# oficial — reimplementar aqui faria os dois documentos divergirem com o tempo. Mesmo precedente
# de hipoteses_utils importando de ia_utils.
from funcoes.relatorios_docx import _montar_rodape, _preparar_documento, _shade

logger = logging.getLogger(__name__)

# Onde os .docx do chat ficam até o usuário baixar. Sai de PASTA_INVESTIGACAO (o bind mount
# ./data do docker-compose), nunca de caminho relativo: o CWD do processo não é contrato.
PASTA_RELATORIOS_IA = Path(PASTA_DADOS) / "relatorios_ia"

_DIAS_MANTER_ARQUIVO = 30   # mesma janela dos .md de "não enviar+exibir prompt"
# Resposta curta não é relatório (é "não encontrei nada", pedido de esclarecimento, recusa...);
# gerar um Word de duas linhas só polui a conversa com um anexo inútil.
_MIN_CHARS_RESPOSTA = 400

CINZA_CABECALHO = "EFEFEF"

INSTRUCAO_RELATORIO = (
    "\n\n## FORMATO DE RELATÓRIO\n"
    "O investigador pediu um RELATÓRIO: a sua resposta deste turno será convertida "
    "automaticamente num documento Word (.docx) para download. Escreva o relatório inteiro na "
    "própria resposta, em markdown, como documento — não como conversa:\n"
    "- Comece por um título com '# ' e organize o texto em seções com '## ' (e '### ' nas "
    "subseções), numa ordem que faça sentido para quem vai ler o documento sozinho, sem a "
    "conversa ao lado.\n"
    "- Nada de saudação, de 'aqui está o relatório' nem de pergunta no fim: o documento começa "
    "no título e termina na última seção.\n"
    "- Use listas com '- ' e tabelas em markdown (| coluna | coluna |) quando organizarem melhor "
    "os dados; não jogue o texto do relatório dentro de bloco de código.\n"
    "- Mantenha os marcadores de evidência ([msg:wpp:ID], [msg:sms:ID], [doc:ID]) junto de cada "
    "afirmação que eles sustentam — eles vão para o documento e são a rastreabilidade da prova.\n"
    "- Separe o que é fato apurado do que é inferência, e feche com as lacunas e as próximas "
    "diligências sugeridas."
)


# =====================================================================
# Reconhecimento do pedido
# =====================================================================

def _sem_acento(texto: str) -> str:
    return unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode("ascii")


# "relatório", "word", "docx" — o que o investigador pode chamar de documento a gerar
_RE_ALVO = re.compile(r"\brelatorios?\b|\bword\b|\bdocx\b")

# verbo de geração (o pedido é "faça um relatório", não "o que diz o relatório")
_RE_VERBO = re.compile(
    r"\b(?:gera|gere|gerar|gerando|elabora|elabore|elaborar|monta|monte|montar|produz|produza|"
    r"produzir|faz|faca|facam|fazer|redige|redija|redigir|escreve|escreva|escrever|cria|crie|"
    r"criar|prepara|prepare|preparar|emite|emita|emitir|consolida|consolide|consolidar|quero|"
    r"queria|gostaria|preciso|precisava|deseja|desejo|me\s+de|me\s+entregue|em\s+formato\s+de|"
    r"em\s+forma\s+de|no\s+formato\s+de)\b"
)

# distância máxima entre o verbo e a palavra-alvo: "gere um relatório completo do inquérito" tem
# poucos caracteres entre os dois; "gostaria que você me entregasse um relatório" tem ~35
_JANELA_VERBO = 80


def pedido_de_relatorio(texto: str) -> bool:
    """True se a mensagem do investigador pede a geração de um relatório/documento Word.

    Regra: uma palavra-alvo ("relatório", "word", "docx") precedida de um verbo de geração dentro
    de uma janela curta, ou a mensagem começando pela própria palavra (o comando telegráfico
    "Relatório...", que o botão 'Gerar relatório' do chat usa). O falso positivo
    custa um .docx a mais no fim da resposta; o falso negativo custa o recurso inteiro — por isso
    a regra é deliberadamente generosa."""
    if not texto:
        return False
    alvo = _sem_acento(texto).lower()
    if alvo.lstrip().startswith("relatorio"):
        return True
    for m in _RE_ALVO.finditer(alvo):
        if _RE_VERBO.search(alvo[max(0, m.start() - _JANELA_VERBO):m.start()]):
            return True
    return False


# =====================================================================
# Markdown -> .docx
# =====================================================================

# negrito | itálico | código | link markdown — na mesma varredura, para não formatar duas vezes
# o mesmo trecho. O marcador de evidência ([msg:wpp:123]) não casa com o link (não vem seguido
# de parêntese) e chega intacto ao documento, que é o que se quer: é a citação da prova.
_RE_INLINE = re.compile(
    r"\*\*(?P<negrito>.+?)\*\*"
    r"|(?<!\*)\*(?P<italico>[^*\n]+)\*(?!\*)"
    r"|`(?P<codigo>[^`\n]+)`"
    r"|\[(?P<rotulo>[^\]\n]+)\]\((?P<url>[^)\n]+)\)"
)

_RE_TITULO = re.compile(r"^(#{1,4})\s+(.*)$")
_RE_ITEM = re.compile(r"^(\s*)[-*+]\s+(.*)$")
_RE_ITEM_NUMERADO = re.compile(r"^(\s*)(\d+)[.)]\s+(.*)$")
_RE_CITACAO = re.compile(r"^\s*>\s?(.*)$")
_RE_LINHA_TABELA = re.compile(r"^\s*\|.*\|\s*$")
_RE_SEPARADOR_TABELA = re.compile(r"^[\s|:-]+$")
_RE_REGUA = re.compile(r"^\s*([-*_])\1{2,}\s*$")


def _escrever_inline(paragrafo, texto: str) -> None:
    """Escreve o texto no parágrafo quebrando as marcações inline em runs formatados.
    De link markdown fica só o rótulo: o .docx é feito para ser impresso e juntado aos autos."""
    pos = 0
    for m in _RE_INLINE.finditer(texto):
        if m.start() > pos:
            paragrafo.add_run(texto[pos:m.start()])
        if m.group("negrito") is not None:
            paragrafo.add_run(m.group("negrito")).bold = True
        elif m.group("italico") is not None:
            paragrafo.add_run(m.group("italico")).italic = True
        elif m.group("codigo") is not None:
            run = paragrafo.add_run(m.group("codigo"))
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        else:
            paragrafo.add_run(m.group("rotulo"))
        pos = m.end()
    if pos < len(texto):
        paragrafo.add_run(texto[pos:])


def _celulas_da_linha(linha: str) -> list[str]:
    return [c.strip() for c in linha.strip().strip("|").split("|")]


def _desenhar_tabela(doc, linhas: list[list[str]]) -> None:
    """Desenha uma tabela markdown já quebrada em células (primeira linha = cabeçalho)."""
    if not linhas:
        return
    colunas = max(len(l) for l in linhas)
    tabela = doc.add_table(rows=0, cols=colunas)
    tabela.style = "Table Grid"
    for idx, celulas in enumerate(linhas):
        linha = tabela.add_row()
        for coluna in range(colunas):
            celula = linha.cells[coluna]
            paragrafo = celula.paragraphs[0]
            _escrever_inline(paragrafo, celulas[coluna] if coluna < len(celulas) else "")
            for run in paragrafo.runs:
                run.font.size = Pt(9)
                if idx == 0:
                    run.bold = True
            if idx == 0:
                _shade(celula, CINZA_CABECALHO)
    doc.add_paragraph()


def escrever_markdown(doc, markdown: str) -> None:
    """Converte o markdown da resposta da IA em parágrafos, listas e tabelas do documento.

    Cobre o mesmo subconjunto do renderizador do chat (web/js/ia.js) — título, negrito/itálico,
    código, lista, lista numerada, tabela, citação e bloco de código —, que é o que os modelos
    realmente usam. O que não for reconhecido vira parágrafo comum: nenhuma linha se perde."""
    linhas = (markdown or "").replace("\r\n", "\n").split("\n")
    tabela: list[list[str]] = []
    em_codigo = False

    def fechar_tabela():
        nonlocal tabela
        if tabela:
            _desenhar_tabela(doc, tabela)
            tabela = []

    for linha in linhas:
        if linha.strip().startswith("```"):
            fechar_tabela()
            em_codigo = not em_codigo
            continue

        if em_codigo:
            paragrafo = doc.add_paragraph()
            paragrafo.paragraph_format.space_after = Pt(0)
            run = paragrafo.add_run(linha)
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            continue

        if _RE_LINHA_TABELA.match(linha):
            if not _RE_SEPARADOR_TABELA.match(linha.strip().strip("|")):
                tabela.append(_celulas_da_linha(linha))
            continue
        fechar_tabela()

        if not linha.strip() or _RE_REGUA.match(linha):
            continue

        titulo = _RE_TITULO.match(linha)
        if titulo:
            cabecalho = doc.add_heading("", level=len(titulo.group(1)))
            _escrever_inline(cabecalho, titulo.group(2).strip())
            continue

        citacao = _RE_CITACAO.match(linha)
        if citacao:
            paragrafo = doc.add_paragraph()
            paragrafo.paragraph_format.left_indent = Cm(1)
            _escrever_inline(paragrafo, citacao.group(1).strip())
            for run in paragrafo.runs:
                run.italic = True
            continue

        item = _RE_ITEM.match(linha)
        if item:
            nivel = min(len(item.group(1)) // 2, 2)
            estilo = "List Bullet" if nivel == 0 else f"List Bullet {nivel + 1}"
            _escrever_inline(doc.add_paragraph(style=estilo), item.group(2).strip())
            continue

        numerado = _RE_ITEM_NUMERADO.match(linha)
        if numerado:
            # o número vai literal, como no chat: o modelo costuma se referir aos itens pelo
            # número ("o indício 3"), e uma lista automática do Word pode renumerar.
            paragrafo = doc.add_paragraph()
            recuo = min(len(numerado.group(1)) // 2, 2) + 1
            paragrafo.paragraph_format.left_indent = Cm(0.75 * recuo)
            _escrever_inline(paragrafo, f"{numerado.group(2)}. {numerado.group(3).strip()}")
            continue

        paragrafo = doc.add_paragraph()
        paragrafo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _escrever_inline(paragrafo, linha.strip())

    fechar_tabela()


def _capa(doc, titulo: str, metadados: list[tuple[str, str]], pergunta: str) -> None:
    """Título, ficha de metadados (de onde saiu, quem pediu, qual modelo) e o pedido literal.
    A resposta da IA vem depois — este bloco é a procedência do documento."""
    paragrafo = doc.add_paragraph()
    paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragrafo.add_run(titulo)
    run.bold = True
    run.font.size = Pt(16)

    tabela = doc.add_table(rows=0, cols=2)
    tabela.style = "Table Grid"
    for rotulo, valor in metadados:
        if not valor:
            continue
        celula_rotulo, celula_valor = tabela.add_row().cells
        run = celula_rotulo.paragraphs[0].add_run(rotulo)
        run.bold = True
        run.font.size = Pt(9)
        _shade(celula_rotulo, CINZA_CABECALHO)
        run = celula_valor.paragraphs[0].add_run(valor)
        run.font.size = Pt(9)

    if pergunta:
        paragrafo = doc.add_paragraph()
        run = paragrafo.add_run("Pedido do investigador: ")
        run.bold = True
        run.font.size = Pt(9)
        run = paragrafo.add_run(pergunta.strip())
        run.font.size = Pt(9)
        run.italic = True

    paragrafo = doc.add_paragraph()
    run = paragrafo.add_run(
        "Documento gerado por inteligência artificial a partir dos dados do sistema. O conteúdo "
        "deve ser conferido nas fontes citadas antes de qualquer uso."
    )
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor.from_string("FF0000")
    run.italic = True

    doc.add_paragraph()


def montar_docx(caminho: Path, titulo: str, pergunta: str, resposta_md: str,
                metadados: list[tuple[str, str]], numero: str = "") -> Path:
    """Monta o .docx da resposta e grava em `caminho` (devolve o próprio caminho)."""
    doc, secao = _preparar_documento()
    _montar_rodape(secao, numero)
    _capa(doc, titulo, metadados, pergunta)
    escrever_markdown(doc, resposta_md)
    caminho.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(caminho))
    return caminho


# =====================================================================
# Persistência (IA_TB05_RELATORIOS_DOCX)
# =====================================================================

def _limpar_relatorios_antigos() -> None:
    """Apaga os .docx com mais de _DIAS_MANTER_ARQUIVO dias e seus registros (limpeza preguiçosa,
    na gravação do próximo). O relatório é regerável; o binário não vale ficar para sempre."""
    marco = datetime.now() - timedelta(days=_DIAS_MANTER_ARQUIVO)
    with get_connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                "SELECT id_relatorio_ia, caminho_arquivo "
                "FROM investigacao.ia_tb05_relatorios_docx WHERE dt_criacao < %s",
                (marco,),
            )
            antigos = cur.fetchall()
            if not antigos:
                return
            for _id, caminho in antigos:
                try:
                    Path(caminho).unlink(missing_ok=True)
                except OSError:
                    pass
            cur.execute(
                "DELETE FROM investigacao.ia_tb05_relatorios_docx WHERE dt_criacao < %s",
                (marco,),
            )
        conn.commit()


def _registrar(id_conversa_ia: int, id_inquerito: int | None,
               id_usuario: int | None, nome: str, caminho: Path) -> int:
    with get_connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                INSERT INTO investigacao.ia_tb05_relatorios_docx
                    (id_conversa_ia, id_inquerito, id_usuario, no_arquivo, caminho_arquivo)
                VALUES (%s, %s, %s, %s, %s)
                RETURNING id_relatorio_ia
                """,
                (id_conversa_ia, id_inquerito, id_usuario, nome, str(caminho)),
            )
            id_relatorio = cur.fetchone()[0]
        conn.commit()
    return id_relatorio


def obter_relatorio_ia(id_relatorio_ia: int) -> dict | None:
    """Registro de um .docx gerado (None se não existir) — usado pelas rotas de download."""
    with get_connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT id_relatorio_ia, id_conversa_ia, id_inquerito, id_usuario,
                       no_arquivo, caminho_arquivo, dt_criacao
                FROM investigacao.ia_tb05_relatorios_docx
                WHERE id_relatorio_ia = %s
                """,
                (id_relatorio_ia,),
            )
            linha = cur.fetchone()
            if not linha:
                return None
            return dict(zip([d.name for d in cur.description], linha))


def _apagar_relatorios_onde(cur, coluna: str, valor: int) -> None:
    """Apaga os .docx casados por COLUNA (arquivo + registro), no cursor de quem está apagando.
    O arquivo sai do disco ANTES da linha: depois do DELETE não há mais como saber qual era, e
    ele ficaria órfão em data/. COLUNA é literal do módulo, nunca vem de fora."""
    cur.execute(
        f"SELECT caminho_arquivo FROM investigacao.ia_tb05_relatorios_docx WHERE {coluna} = %s",
        (valor,),
    )
    for (caminho,) in cur.fetchall():
        try:
            Path(caminho).unlink(missing_ok=True)
        except OSError:
            pass
    cur.execute(
        f"DELETE FROM investigacao.ia_tb05_relatorios_docx WHERE {coluna} = %s",
        (valor,),
    )


def apagar_relatorios_da_conversa(cur, id_conversa_ia: int) -> None:
    """Apaga os .docx de uma conversa (arquivo + registro), no cursor de quem apaga a conversa."""
    _apagar_relatorios_onde(cur, "id_conversa_ia", id_conversa_ia)


def apagar_relatorios_de_inquerito(cur, id_inquerito: int) -> None:
    """Todos os .docx do inquérito de uma vez — a FK243 para INQ_TB01 não tem CASCADE, então sem
    isto o inquérito inteiro não sai (ID_CONVERSA_IA não tem FK: apagar a conversa não basta)."""
    _apagar_relatorios_onde(cur, "id_inquerito", id_inquerito)


# =====================================================================
# Ponto de entrada do chat
# =====================================================================

_RE_NOME_INVALIDO = re.compile(r"[^0-9A-Za-z._-]+")


def _nome_arquivo(titulo: str, id_conversa_ia: int) -> str:
    base = _RE_NOME_INVALIDO.sub("_", _sem_acento(titulo or "").strip()).strip("_")[:60]
    return f"relatorio_ia_{base or id_conversa_ia}_{datetime.now():%Y%m%d_%H%M%S}.docx"


def _sem_marcacao(texto: str) -> str:
    return _RE_INLINE.sub(
        lambda m: m.group("negrito") or m.group("italico") or m.group("codigo") or m.group("rotulo") or "",
        texto,
    ).strip()


def _titulo_e_corpo(resposta_md: str, titulo_conversa: str) -> tuple[str, str]:
    """(título do documento, corpo sem a linha do título).

    O título é o primeiro '# '/'## ' que o modelo escreveu — a instrução do prompt pede que ele
    abra o relatório assim. Essa linha sai do corpo: senão o título apareceria duas vezes, na capa
    e como primeira seção. Sem título no markdown, vale o da conversa e o corpo vai inteiro."""
    linhas = (resposta_md or "").split("\n")
    for i, linha in enumerate(linhas[:20]):
        titulo = _RE_TITULO.match(linha)
        if not titulo:
            continue
        if len(titulo.group(1)) > 2:  # '### ' já é subseção: é conteúdo, não o título do documento
            break
        texto = _sem_marcacao(titulo.group(2))
        if texto:
            return texto[:200], "\n".join(linhas[:i] + linhas[i + 1:])
        break
    return (titulo_conversa or "Relatório").strip()[:200], resposta_md


def _nome_usuario(id_usuario: int | None) -> str:
    if not id_usuario:
        return ""
    with get_connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                "SELECT no_usuario FROM seguranca.usr_tb01_usuarios WHERE id_usuario = %s",
                (id_usuario,),
            )
            linha = cur.fetchone()
    return (linha[0] or "") if linha else ""


def _dono(id_inquerito: int | None) -> tuple[str, str]:
    """(rótulo, valor) da origem do relatório — número do inquérito."""
    with get_connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                "SELECT nu_inquerito FROM investigacao.inq_tb01_inqueritos WHERE id_inquerito = %s",
                (id_inquerito,),
            )
            linha = cur.fetchone()
            return "Inquérito", (linha[0] if linha and linha[0] else str(id_inquerito))


def anexar_relatorio(resposta: str, pergunta: str, titulo_conversa: str, id_conversa_ia: int,
                     id_usuario: int | None, modelo: str | None, rota_download: str,
                     id_inquerito: int | None = None) -> str:
    """Gera o .docx desta resposta e devolve a resposta com o link de download no fim.

    `rota_download` é o prefixo da rota que serve o arquivo ('/api/investigacao/ia/relatorios'):
    o chat confere o acesso pela própria conversa.

    Nunca levanta: neste ponto a resposta da IA já existe e será gravada de qualquer jeito — um
    erro ao montar o Word vira log e a conversa segue sem o anexo."""
    if not resposta or len(resposta.strip()) < _MIN_CHARS_RESPOSTA:
        return resposta
    try:
        titulo, corpo = _titulo_e_corpo(resposta, titulo_conversa)
        rotulo_dono, valor_dono = _dono(id_inquerito)
        metadados = [
            (rotulo_dono, valor_dono),
            ("Gerado em", datetime.now().strftime("%d/%m/%Y %H:%M")),
            ("Investigador", _nome_usuario(id_usuario)),
            ("Modelo de IA", modelo or ""),
        ]
        nome = _nome_arquivo(titulo, id_conversa_ia)
        caminho = montar_docx(
            PASTA_RELATORIOS_IA / nome, titulo, pergunta, corpo, metadados,
            numero=valor_dono if rotulo_dono == "Inquérito" else "",
        )
        _limpar_relatorios_antigos()
        id_relatorio = _registrar(id_conversa_ia, id_inquerito, id_usuario, nome, caminho)
    except Exception:
        logger.exception("Falha ao gerar o .docx do relatório da conversa %s", id_conversa_ia)
        return resposta
    return (
        f"{resposta}\n\n"
        f"📄 [Baixar este relatório em Word (.docx)]({rota_download}/{id_relatorio}/download)"
    )
